from docx import Document
import flask
from flask import request, jsonify, send_file
from fpdf import FPDF 
import os

app = flask.Flask(__name__)
app.config["DEBUG"] = True


@app.route('/', methods=['GET'])
def home():
    if 'Equipement' in request.args:
        return find(request.args['Equipement'])
    else:
        return "Error: No Equipement field provided. Please specify an Equipement."



def find(title):
    doc = Document("mach.docx")

    fulldoc = []
    grouped = []


    for para in doc.paragraphs:
        fulldoc.append(para.text)

    for i in range(0,len(fulldoc)):
        if len(fulldoc[i]) < 20:
            for par in range(i+1,len(fulldoc)):
                if len(fulldoc[par]) > 20:
                    fulldoc[i] += "\n" + fulldoc[par]
                else:
                    break


    for par in fulldoc:
        if  ("\n" in par):
            grouped.append(par)
            


    

    for par in grouped:
        if title.lower() in par.lower():
            pdf = FPDF() 
  
            pdf.add_page() 
            
            pdf.set_font("Arial",'B', size = 15) 
            pdf.cell(200, 20, txt = par.split("\n")[0], ln = 1, align = 'C')
            pdf.image("circle.png",10,10,25,20)

            pdf.cell(200, 10, txt = "\n", ln = 1, align = 'C') 
            pdf.set_font("Times", size = 12) 
            for l in par.split("\n")[1:]:
                pdf.multi_cell( h=5.0, align='L', w=0, txt=l, border=0)
                pdf.cell(200, 5, txt = "\n", ln = 1)
            # save the pdf with name .pdf 
            pdf.output("GFG.pdf")
            return send_file(os.path.join( "GFG.pdf"))
            break
        else:
            return "Machine not found"

app.run()
